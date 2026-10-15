"""R6+1 — note BUG-123 partial fix + flip BUG-137 (co-fixed by BUG-136)."""
import openpyxl

PATH = "docs/audits/Multiship_Bug_List_2026-10-14_R6.xlsx"
TODAY = "2026-10-15"

wb = openpyxl.load_workbook(PATH)
ws = wb["Bugs"]
hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}

for row in ws.iter_rows(min_row=2):
    bid = row[0].value
    if bid == "BUG-123":
        prior = ws.cell(row=row[0].row, column=hdr["Developer comments"]).value or ""
        add = "R6+1 (8ab23db5, " + TODAY + "): CountrySelect now carries ArrowDown/Up/Enter/Escape keyboard support + aria-activedescendant / role=combobox / role=option so the shared country picker is reachable from the keyboard. HS code combobox landed earlier. Broader dialog focus/escape audit stays Deferred - that is a cross-cutting sweep across every Settings modal."
        combined = (prior + "\n" + add) if prior and add not in prior else add
        ws.cell(row=row[0].row, column=hdr["Developer comments"]).value = combined
        ws.cell(row=row[0].row, column=hdr["Fixed in build"]).value = TODAY
        print("BUG-123 partial-fix note appended (stays Deferred).")
    elif bid == "BUG-137":
        prior = ws.cell(row=row[0].row, column=hdr["Developer comments"]).value or ""
        add = ("R6+1 (" + TODAY + "): FedExConnector builds the shipper block from "
               "request.getShipperXxx() verbatim - no tenant-settings fallback. "
               "BUG-136 fix (c39e9e6e) stops the FE from pre-filling the warehouse "
               "address in Return mode, so every NEW return label will carry the "
               "customer's address on SHIP FROM. Historical ZPL for #934194 remains "
               "unchanged (stored label is immutable).")
        combined = (prior + "\n" + add) if prior and add not in prior else add
        ws.cell(row=row[0].row, column=hdr["Dev status"]).value = "Fixed"
        ws.cell(row=row[0].row, column=hdr["Developer comments"]).value = combined
        ws.cell(row=row[0].row, column=hdr["Fixed in build"]).value = TODAY
        print("BUG-137 flipped Fixed.")

wb.save(PATH)
