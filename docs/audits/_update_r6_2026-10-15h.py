"""R6+1 — flip BUG-129 + BUG-136 Fixed after c39e9e6e."""
import openpyxl

PATH = "docs/audits/Multiship_Bug_List_2026-10-14_R6.xlsx"
TODAY = "2026-10-15"

updates = {
    "BUG-129": "R6+1 (c39e9e6e, " + TODAY + "): AuditLogPage dropdowns extended to cover every AuditService entity + action constant (IMPORT_BATCH + LABEL_REGENERATED / LABEL_VOID_REFUSED / CARRIER_REJECTED / IMPORT_SAVED / IMPORT_GENERATED / LOGIN_FAILED). Hardcoded against the backend constants rather than DB-driven - swap to a /api/v1/audit-log/filter-values endpoint if the enum ever grows.",
    "BUG-136": "R6+1 (c39e9e6e, " + TODAY + "): switching SHIPMENT -> RETURN now clears the sender block when it still matches the warehouse address exactly, so 'Return from * customer' opens blank for the operator to fill. Anything typed is preserved.",
}

wb = openpyxl.load_workbook(PATH)
ws = wb["Bugs"]
hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}

for row in ws.iter_rows(min_row=2):
    bid = row[0].value
    if bid in updates:
        prior = ws.cell(row=row[0].row, column=hdr["Developer comments"]).value or ""
        add = updates[bid]
        combined = f"{prior}\n{add}" if prior and add not in prior else add
        ws.cell(row=row[0].row, column=hdr["Dev status"]).value = "Fixed"
        ws.cell(row=row[0].row, column=hdr["Developer comments"]).value = combined
        ws.cell(row=row[0].row, column=hdr["Fixed in build"]).value = TODAY
        print(f"{bid} Fixed.")

wb.save(PATH)
