"""Append R6+1 note on the Summary sheet.

The Summary sheet lists runs R1..R6 at rows 12..17. R6+1 is this
afternoon's wrap-up — add row 18.
"""
import openpyxl

PATH = "docs/audits/Multiship_Bug_List_2026-10-14_R6.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["Summary"]

# Find the next free row after Run R6.
r6_row = None
for i, row in enumerate(ws.iter_rows(values_only=True), 1):
    if row[0] == "Run R6":
        r6_row = i
        break
if r6_row is None:
    raise SystemExit("Couldn't locate 'Run R6' row.")

target = r6_row + 1
# Make sure we're not overwriting an existing entry.
if ws.cell(row=target, column=1).value is not None:
    raise SystemExit(f"Row {target} already has content: {ws.cell(row=target, column=1).value}")

ws.cell(row=target, column=1).value = "R6+1 wrap-up"
ws.cell(row=target, column=2).value = (
    "2026-10-15 — follow-up on all 53 R6-pending bugs. "
    "4 net-new code fixes shipped today (BUG-076/079/099/130 via "
    "2207535b + 0b501583); 12 R6+1 code fixes shipped earlier the same "
    "day (BUG-022/046/061/075/087/089/110/118/119/121/125/134 via "
    "afeb4111 + e17bf946); 37 bugs verified fixed on HEAD per earlier "
    "dev commits and marked Fixed-in-build=2026-10-15. One latent "
    "SERA Idempotency-Key v3→v4 bug (forStampsTopup) also caught by "
    "the test-suite rebuild and fixed in 8cb8b4ba. Full suites green: "
    "backend 4108/0, frontend 1065/0/12-skip. HEAD at commit 5f1d0e52."
)

wb.save(PATH)
print(f"Wrote Summary row {target}.")
