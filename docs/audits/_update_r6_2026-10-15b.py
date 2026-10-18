"""Add R6+1 notes for the test-suite-catchup commits of 2026-10-15.

afeb4111/e17bf946 were the primary R6+1 fix commits (already recorded).
8cb8b4ba (backend) + 46d6db84 (frontend) are test-suite catch-up — the
suites were asserting pre-fix behaviour and had drifted. One real latent
bug was caught: forStampsTopup was returning a v3 UUID; SERA rejects
non-v4 keys with error 800001. Routed through stampsKey() so topup is
now v4 like every other forStamps* helper.

For bugs already marked Fixed in a prior tracker pass, we append a one-
line note so the audit trail shows the test now locks in the behaviour.
"""
import openpyxl

PATH = "docs/audits/Multiship_Bug_List_2026-10-14_R6.xlsx"
TODAY = "2026-10-15"
BE_SHA = "8cb8b4ba"
FE_SHA = "46d6db84"

notes = {
    # Primary — latent bug caught by test suite.
    "BUG-118": (
        f"Test-suite fix ({BE_SHA}, {TODAY}): also caught a latent bug — "
        "forStampsTopup returned a v3 UUID (nameUUIDFromBytes), which "
        "SERA rejects with error 800001 'Idempotency-Key invalid'. "
        "Topup now routes through stampsKey() like every other "
        "forStamps* helper so the version bits are stamped v4. "
        "Backfilled test: IdempotencyKeysStampsSeraTest.everyStampsKeyIsAStableV4Uuid."),

    # Pure test catch-ups (behaviour was already shipped in R6/R6+1).
    "BUG-022": (f"Test fix ({FE_SHA}, {TODAY}): — no code change, just test lock-in."),
    "BUG-029": (f"Test fix ({FE_SHA}, {TODAY}): tests now assert UPS maxLength 10 + "
                "DHL maxLength 12 to match the 6-10 / 9-12 widening."),
    "BUG-057": (f"Test fix ({FE_SHA}, {TODAY}): — no code change, just test lock-in."),
    "BUG-061": (f"Test fix ({FE_SHA}, {TODAY}): SystemChannelSection test no longer "
                "asserts the scrubbed raw TENANT_CHANNEL_NOT_ENABLED code is rendered."),
    "BUG-109": (f"Test fix ({FE_SHA}, {TODAY}): fetchPreviewObjectUrl now asserts the "
                "second-arg template id (previewContext.id) that landed with 0fad4a45."),
    "BUG-138": (f"Test fix ({FE_SHA}, {TODAY}): importer-broker fixture now seeds "
                "non-blank clientId+clientSecret so the carriersStepComplete guard "
                "lets Submit through."),
}

wb = openpyxl.load_workbook(PATH)
ws = wb["Bugs"]
hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}

updated = []
for row in ws.iter_rows(min_row=2):
    bid = row[0].value
    if bid in notes:
        prior = ws.cell(row=row[0].row, column=hdr["Developer comments"]).value or ""
        add = notes[bid]
        if add not in prior:
            ws.cell(row=row[0].row, column=hdr["Developer comments"]).value = f"{prior}\n{add}" if prior else add
            updated.append(bid)

wb.save(PATH)
print(f"Appended {len(updated)}: {', '.join(updated)}")
