"""R6+1 BUG-128 — flip Deferred → Fixed after sprint 2 (f14bcf08).

Sprint 1 (f46b6f01) wired AuditService into 7 services. Sprint 2
(f14bcf08) covers the remaining 4 Settings surfaces the tester named:
Printers, Cutoffs & Holidays, Address book, Output Destinations.
"""
import openpyxl

PATH = "docs/audits/Multiship_Bug_List_2026-10-14_R6.xlsx"
TODAY = "2026-10-15"

wb = openpyxl.load_workbook(PATH)
ws = wb["Bugs"]
hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}

for row in ws.iter_rows(min_row=2):
    if row[0].value == "BUG-128":
        prior = ws.cell(row=row[0].row, column=hdr["Developer comments"]).value or ""
        add = (f"R6+1 sprint 2 (f14bcf08, {TODAY}): AuditService wired into "
               "PrinterService (Settings > Printers CRUD), CutoffAdminController "
               "(cutoff rules + holidays), SavedRecipientServiceImpl (Address "
               "book) and OutputDestinationAdminService (output destinations). "
               "Pickups already recorded via CarrierEodLogger. "
               "38/38 related unit tests green.")
        combined = f"{prior}\n{add}" if prior and add not in prior else add
        ws.cell(row=row[0].row, column=hdr["Dev status"]).value = "Fixed"
        ws.cell(row=row[0].row, column=hdr["Developer comments"]).value = combined
        ws.cell(row=row[0].row, column=hdr["Fixed in build"]).value = TODAY
        print("BUG-128 flipped to Fixed.")
        break

wb.save(PATH)
