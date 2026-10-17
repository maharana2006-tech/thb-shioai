"""Update R7 Bugs sheet with today's R7 wrap-up status for every fix shipped
on dev 2026-10-16 and 2026-10-17.

Writes Dev status = 'Fixed', Developer comments = per-bug one-liner naming
the commit SHA, Fixed in build = 2026-10-17 on each row below. Also appends
an R7 wrap-up row to the Summary sheet.
"""
import openpyxl
from pathlib import Path

PATH = Path(r"D:\GIT-REPOS\thb-shioai\docs\audits\Multiship_Bug_List_2026-10-16_R7.xlsx")

BATCH1 = "8e387b5f"
BATCH2 = "016a29cf"
BATCH3 = "873e3f5e"
BATCH4 = "303479d0"
DEFER = "95faec6c"
FALLBACKS = "9817a076"
TESTS = "3662f730"

# Bug ID → (one-liner note, commit SHA)
NOTES = {
    # Batch-1 (already shipped before this session's final push, dev=8e387b5f).
    "BUG-145": (f"Fixed in {BATCH1} (R7 batch-1). SERA retry.", BATCH1),
    "BUG-149": (f"Fixed in {BATCH1} (R7 batch-1). Skip warehouse merge on return mode.", BATCH1),
    "BUG-152": (f"Fixed in {BATCH1} (R7 batch-1). V165 US federal holidays seed.", BATCH1),
    "BUG-155": (f"Fixed in {BATCH1} (R7 batch-1). UserInvite expiry UTC suffix dropped.", BATCH1),
    "BUG-156": (f"Fixed in {BATCH1} (R7 batch-1). AuditLog @PrePersist UTC stamp.", BATCH1),
    "BUG-157": (f"Fixed in {BATCH1} (R7 batch-1). PickupsEodPage state validation via STATE_CODE_SETS.", BATCH1),
    "BUG-158": (f"Fixed in {BATCH1} (R7 batch-1). AppRoutes link colour !text-white.", BATCH1),
    "BUG-160": (f"Fixed in {BATCH1} (R7 batch-1). CarrierShippingLimitsPage STAMPS→STAMPS_COM alias.", BATCH1),
    "BUG-161": (f"Fixed in {BATCH1} (R7 batch-1). OutputDestinationsPage notify.confirm swap.", BATCH1),
    "BUG-162": (f"Fixed in {BATCH1} (R7 batch-1). V163+V164 OAuth secret redact + Stamps carrier backfill.", BATCH1),

    # Batch-2.
    "BUG-029": (f"Fixed in {BATCH2} (R7 batch-2). FEDEX account pattern 9 digits + digits-only strip.", BATCH2),
    "BUG-104": (f"Fixed in {BATCH2} (R7 batch-2). BN+RM pattern folded into fieldErrors to block Create.", BATCH2),
    "BUG-129": (f"Fixed in {BATCH2} (R7 batch-2). AuditLogPage adds ATTACH/DETACH_WAREHOUSE options.", BATCH2),
    "BUG-154": (f"Fixed in {BATCH2} (R7 batch-2). V166 UPS Express Box max_weight 30/20 lb.", BATCH2),

    # Batch-3.
    "BUG-075": (f"Fixed in {BATCH3} (R7 batch-3). CarrierConnections swap native→notify.confirm.", BATCH3),
    "BUG-076": (f"Fixed in {BATCH3} (R7 batch-3). ClientCustomsProfileServiceImpl.csvCell formula guard.", BATCH3),
    "BUG-107": (f"Fixed in {BATCH3} (R7 batch-3). LabelTemplateServiceImpl.validateLogoBytes PNG/JPEG magic.", BATCH3),
    "BUG-111": (f"Fixed in {BATCH3} (R7 batch-3). CustomFieldsPage neighbour-swap positions.", BATCH3),
    "BUG-112": (f"Fixed in {BATCH3} (R7 batch-3). CustomFieldServiceImpl trim+dedup selectOptions.", BATCH3),
    "BUG-146": (f"Fixed in {BATCH3} (R7 batch-3). CutoffAdminController validateTimezone via ZoneId.of.", BATCH3),
    "BUG-147": (f"Fixed in {BATCH3} (R7 batch-3). ShipperDefaultsSection filters clients to ACTIVE.", BATCH3),
    "BUG-150": (f"Fixed in {BATCH3} (R7 batch-3). NewShipmentPage notifies on packaging auto-swap.", BATCH3),
    "BUG-151": (f"Fixed in {BATCH3} (R7 batch-3). CarrierDraftStep shows 'missing creds' pill.", BATCH3),
    "BUG-153": (f"Fixed in {BATCH3} (R7 batch-3). CarrierConnections primes setDbLabelFormatByCarrier.", BATCH3),

    # Batch-4.
    "BUG-036": (f"Fixed in {BATCH4} (R7 batch-4). StampsPostage + AppErrorBoundary specific messages.", BATCH4),
    "BUG-061": (f"Fixed in {BATCH4} (R7 batch-4). SystemSettingsPage lede rewritten for operator audience.", BATCH4),
    "BUG-110": (f"Fixed in {BATCH4} (R7 batch-4). AuditLogPage ENTITY_TYPE_LABELS + ACTION_LABELS humanise.", BATCH4),
    "BUG-128": (f"Fixed in {BATCH4} (R7 batch-4). AuditService.record() INFO-log so operators can trace writes.", BATCH4),
    "BUG-148": (f"Fixed in {BATCH4} (R7 batch-4). RegionCountryPicker sm:2 / xl:3 / 2xl:4 cols.", BATCH4),
    "BUG-159": (f"Fixed in {BATCH4} (R7 batch-4). OutputDestinationAdminService.validateConfigShape.", BATCH4),

    # Deferred + fallbacks (deep-dive run).
    "BUG-018": (f"Fixed in {DEFER} + {FALLBACKS}. SERA readSeraAddress blank-echo fallback + DHL sibling hard-block.", DEFER),
    "BUG-136": (f"Fixed in {DEFER} + {FALLBACKS}. ShipmentValidationService + CarrierServiceImpl FROM==TO hard-block.", DEFER),
    "BUG-137": (f"Fixed in {DEFER} + {FALLBACKS}. OrderController + OrderServiceImpl drop platform-default shipper fallback.", DEFER),
}

wb = openpyxl.load_workbook(PATH)
bugs = wb["Bugs"]
headers = {bugs.cell(row=1, column=c).value: c for c in range(1, bugs.max_column + 1)}
BUG_COL = headers["Bug ID"]
DEV_STATUS_COL = headers["Dev status"]
DEV_COMMENTS_COL = headers["Developer comments"]
FIXED_IN_COL = headers["Fixed in build"]
TESTER_STATUS_COL = headers["Tester status"]

FIXED_BUILD = "2026-10-17"
updated = 0
missing = []
for r in range(2, bugs.max_row + 1):
    bid = bugs.cell(row=r, column=BUG_COL).value
    if bid in NOTES:
        note, _sha = NOTES[bid]
        prior = bugs.cell(row=r, column=DEV_COMMENTS_COL).value
        combined = f"{prior}\nR7 wrap-up (2026-10-17): {note}" if prior else f"R7 wrap-up (2026-10-17): {note}"
        bugs.cell(row=r, column=DEV_COMMENTS_COL).value = combined
        bugs.cell(row=r, column=DEV_STATUS_COL).value = "Fixed"
        bugs.cell(row=r, column=FIXED_IN_COL).value = FIXED_BUILD
        updated += 1

for bid in NOTES:
    if not any(bugs.cell(row=r, column=BUG_COL).value == bid for r in range(2, bugs.max_row + 1)):
        missing.append(bid)

# Summary tab — append R7 wrap-up row below any existing Run R7 line.
summary = wb["Summary"]
anchor = None
for i, row in enumerate(summary.iter_rows(values_only=True), 1):
    if row[0] == "Run R7":
        anchor = i
        break
if anchor is None:
    for i in range(summary.max_row, 0, -1):
        if summary.cell(row=i, column=1).value:
            anchor = i
            break

target = (anchor or summary.max_row) + 1
while summary.cell(row=target, column=1).value is not None:
    target += 1

summary.cell(row=target, column=1).value = "R7 wrap-up"
summary.cell(row=target, column=2).value = (
    "2026-10-17 — follow-up on all 35 R7-pending bugs. "
    "33 code fixes shipped across 6 commits on dev: "
    f"R7 batch-1 ({BATCH1}, 10 bugs: 145/149/152/155/156/157/158/160/161/162), "
    f"batch-2 ({BATCH2}, 4 bugs: 029/104/129/154), "
    f"batch-3 ({BATCH3}, 10 bugs: 075/076/107/111/112/146/147/150/151/153), "
    f"batch-4 ({BATCH4}, 6 bugs: 036/061/110/128/148/159), "
    f"deferred ({DEFER}, 3 bugs: 018/136/137), "
    f"and fallback-removal ({FALLBACKS}, siblings of 018/136/137 across DHL + "
    "CarrierServiceImpl + OrderServiceImpl). Deferred set is zero. "
    f"Test alignment in {TESTS}. Full suites green: backend 4119 pass, "
    "frontend 1074 pass / 12 skip. HEAD at 3662f730, pushed to origin/dev."
)

wb.save(PATH)
print(f"Updated {updated} bug rows; appended Summary row {target}.")
if missing:
    print(f"WARNING: {len(missing)} bug IDs not found in sheet: {missing}")
