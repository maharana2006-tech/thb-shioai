"""Append dev status to a bug row in Multiship_Bug_List_Settings_2026-10-12.xlsx.

Usage:
  python _update_bug.py BUG-001 "Fixed on dev" "<sha>" "<yyyy-mm-dd>" "<optional dev comments>"
Status options:  "Fixed on dev" | "Deferred" | "Query answered" | "Won't fix (by design)" | "Needs product input"
"""
import sys
from pathlib import Path
import openpyxl

if len(sys.argv) < 5:
    print(__doc__, file=sys.stderr); sys.exit(2)
bug_id, status, sha, fixed_in = sys.argv[1:5]
comments = sys.argv[5] if len(sys.argv) >= 6 else ""

xlsx = Path(__file__).parent / "Multiship_Bug_List_Settings_2026-10-12.xlsx"
wb = openpyxl.load_workbook(xlsx)
ws = wb["Bugs"]
# Header: Bug ID(1) ... Dev status(26) Assigned to(27) Developer comments(28) Fixed in build(29)
dev_status_col, dev_comments_col, fixed_in_col = 26, 28, 29
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value == bug_id:
        ws.cell(r, dev_status_col).value = status
        existing = ws.cell(r, dev_comments_col).value or ""
        new_comment = f"[{sha[:8]}] {comments}" if comments else f"[{sha[:8]}]"
        ws.cell(r, dev_comments_col).value = (existing + ("; " if existing else "") + new_comment).strip("; ")
        ws.cell(r, fixed_in_col).value = fixed_in
        wb.save(xlsx)
        print(f"OK {bug_id}: {status} / {sha[:8]} / {fixed_in}")
        sys.exit(0)
print(f"ERROR: {bug_id} not found in {xlsx}", file=sys.stderr); sys.exit(1)
