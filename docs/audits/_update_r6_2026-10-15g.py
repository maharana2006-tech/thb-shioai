"""R6+1 — flip 3 Deferred bugs after sweep.

BUG-085 + BUG-086 already landed per earlier dev commits (tracker was
stale). BUG-090 fixed today via c96e743f.

Deferred count after this: 13 -> 10.
"""
import openpyxl

PATH = "docs/audits/Multiship_Bug_List_2026-10-14_R6.xlsx"
TODAY = "2026-10-15"

updates = {
    "BUG-085": "Already on HEAD — AdminUserInviteController calls adminUserService.recordInviteResent / recordInviteRevoked which writes INVITE_RESENT / INVITE_REVOKED rows to user_admin_audit so Recent admin actions picks them up.",
    "BUG-086": "Already on HEAD — MailProviderFormModal uses autoComplete='new-password' on every secret-typed field so Chrome stops dumping the Multiship sign-in into them.",
    "BUG-090": "R6+1 (c96e743f, " + TODAY + "): ShippingCatalogPage lifts origin + onOriginChange and passes to both ShippingServicesPage and PackagesPage, so the sub-tabs now share the filter the page description promised.",
}

wb = openpyxl.load_workbook(PATH)
ws = wb["Bugs"]
hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}

for row in ws.iter_rows(min_row=2):
    bid = row[0].value
    if bid in updates:
        prior = ws.cell(row=row[0].row, column=hdr["Developer comments"]).value or ""
        add = f"R6+1 ({TODAY}): {updates[bid]}"
        combined = f"{prior}\n{add}" if prior and add not in prior else add
        ws.cell(row=row[0].row, column=hdr["Dev status"]).value = "Fixed"
        ws.cell(row=row[0].row, column=hdr["Developer comments"]).value = combined
        ws.cell(row=row[0].row, column=hdr["Fixed in build"]).value = TODAY
        print(f"{bid} flipped Fixed.")

wb.save(PATH)
