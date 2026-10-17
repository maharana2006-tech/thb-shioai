"""Mark the 6 remaining R7 reopened/deferred bugs Fixed in build 2026-10-17.

Appends an R7 wrap-up-2 note per bug (previous R7 wrap-up on dev=3662f730
covered 33 bugs; these are the 6 Reopened/Deferred items that followed).
Also adds a Summary row below "R7 wrap-up".
"""
import openpyxl
from pathlib import Path

PATH = Path(r"D:\GIT-REPOS\thb-shioai\docs\audits\Multiship_Bug_List_2026-10-16_R7.xlsx")

DEFERRED6 = "f96b539f"
TESTS = "dc031230"

NOTES = {
    "BUG-007": (f"Fixed in {DEFERRED6}. LabelArtifactResolver caches fetched URL bytes back inline (base64) on first read; subsequent /label fetches skip the FedEx cache round-trip + survive URL expiry."),
    "BUG-017": (f"Fixed in {DEFERRED6}. V167 extends V150's weight-backfill to orders without label_package rows by summing packages_json weights into label_batch.weight."),
    "BUG-070": (f"Fixed in {DEFERRED6}. ApiResponse.timestamp @Builder.Default's to now(); GlobalExceptionHandler flipped 10 x status('error')→status('ERROR'); handleTypeMismatch populates errors.field; new handleNoResource (NoResourceFoundException) + ErrorCode.ROUTE_NOT_FOUND close the 3-envelope split."),
    "BUG-088": (f"Fixed in {DEFERRED6}. OrderTrackingRepository.aggregateUsageByAccount coalesces label_generated_at with updated_at/created_at so legacy rows surface a real lastUsedAt."),
    "BUG-097": (f"Fixed in {DEFERRED6}. V169 re-applies V160's FedEx Pak / UPS Express Pak 3→5 lb fix (V160 shipped but DB didn't converge). Idempotent via max_weight < 5."),
    "BUG-123": (f"Fixed in {DEFERRED6}. Four Settings dialogs migrated to useModalDismiss (focus trap + Escape + focus-restore): InviteUserDialog, TrackingTimelineModal, PrinterEditor, ShippingServicesPage Allowed packages."),
}

wb = openpyxl.load_workbook(PATH)
bugs = wb["Bugs"]
headers = {bugs.cell(row=1, column=c).value: c for c in range(1, bugs.max_column + 1)}
BUG_COL = headers["Bug ID"]
DEV_STATUS_COL = headers["Dev status"]
DEV_COMMENTS_COL = headers["Developer comments"]
FIXED_IN_COL = headers["Fixed in build"]

FIXED_BUILD = "2026-10-17"
updated = 0
missing = []
for r in range(2, bugs.max_row + 1):
    bid = bugs.cell(row=r, column=BUG_COL).value
    if bid in NOTES:
        prior = bugs.cell(row=r, column=DEV_COMMENTS_COL).value
        line = f"R7 wrap-up-2 (2026-10-17): {NOTES[bid]}"
        combined = f"{prior}\n{line}" if prior else line
        bugs.cell(row=r, column=DEV_COMMENTS_COL).value = combined
        bugs.cell(row=r, column=DEV_STATUS_COL).value = "Fixed"
        bugs.cell(row=r, column=FIXED_IN_COL).value = FIXED_BUILD
        updated += 1

for bid in NOTES:
    if not any(bugs.cell(row=r, column=BUG_COL).value == bid for r in range(2, bugs.max_row + 1)):
        missing.append(bid)

summary = wb["Summary"]
anchor = None
for i, row in enumerate(summary.iter_rows(values_only=True), 1):
    if row[0] == "R7 wrap-up":
        anchor = i
        break
if anchor is None:
    anchor = summary.max_row

target = anchor + 1
while summary.cell(row=target, column=1).value is not None:
    target += 1

summary.cell(row=target, column=1).value = "R7 wrap-up-2"
summary.cell(row=target, column=2).value = (
    "2026-10-17 — deep-dive on all 6 R7 Reopened/Deferred bugs that followed the "
    "33-bug R7 wrap-up: BUG-007 (FedEx cache URL — lazy-cache bytes inline), "
    "BUG-017 (ERROR-order weight backfill via V167), BUG-070 (API error envelope — "
    "@Builder.Default timestamp + status casing + NoResourceFoundException handler + "
    "ROUTE_NOT_FOUND), BUG-088 (lastUsedAt COALESCE), BUG-097 (V169 re-apply Pak 5 lb), "
    f"BUG-123 (4 Settings dialogs → useModalDismiss). All in {DEFERRED6}; test "
    f"alignment in {TESTS}. Also fixed a duplicate-V163 collision by renaming "
    "V163__redact_historical_oauth_secrets.sql → V168. Deferred set is now empty. "
    "Full suites green: backend 4119 pass, frontend 1074 pass / 12 skip. HEAD at "
    "dc031230, pushed to origin/dev."
)

wb.save(PATH)
print(f"Updated {updated} bug rows; appended Summary row {target}.")
if missing:
    print(f"WARNING: {len(missing)} bug IDs not found in sheet: {missing}")
